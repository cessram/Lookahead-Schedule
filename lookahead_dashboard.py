"""
Construction Lookahead Dashboard v4 (2 / 3 / 4 / 6 week) - Sleek Gantt + Bird logo
-----------------------------------------------------------------------
- Upload the Bird "6wk Schedule" lookahead template (.xlsx / .xlsm)
- Pick a 2, 3, 4 or 6-week window -> sleek Gantt chart
- Enter % complete per activity -> shown directly on the Gantt bars
- Download the lookahead in the SAME Excel format (progress written to
  STATUS / % COMPLETE only where entered; blanks are left as they were)

Run:
    pip install streamlit pandas plotly openpyxl
    streamlit run lookahead_dashboard_v4.py

Logo:
    Save Bird.png in the SAME folder as this script (or set LOGO_PATH below).
"""

import io
import re
from pathlib import Path
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

# ============================================================== look & feel
st.set_page_config(page_title="Lookahead Schedule Dashboard",
                   page_icon="🏗️", layout="wide")

FONT = "Segoe UI, Inter, Helvetica Neue, Arial, sans-serif"
TEXT = "#1F2937"
MUTED = "#9CA3AF"
GRID = "#F1F3F6"
ACCENT = "#E8833A"          # softened Bird orange
TODAY = "#EF4444"

# (planned / light, completed / dark) per trade-section - modern muted tones
PALETTE = [
    ("#C7D7F5", "#4F7CDB"),  # indigo blue
    ("#BFE0EE", "#2B8CB8"),  # sky blue
    ("#FBDCC3", "#E8833A"),  # Bird orange
    ("#E1D5F5", "#8B5CF6"),  # violet
    ("#F8CFD6", "#E05A74"),  # rose
    ("#C4ECEC", "#14A3A3"),  # teal
    ("#F3E6B5", "#C99A12"),  # amber
    ("#D9DEE6", "#5B6B82"),  # slate
]

# Bird scope is always dark green (planned = dark green, completed = deeper green)
BIRD_COLORS = ("#2E7D4F", "#14532D")

# Logo: Bird.png next to this script
LOGO_PATH = Path(__file__).parent / "Bird.png"


def is_bird_section(name):
    return "bird" in str(name).lower()


# Excel colours for the export
XL_PLANNED = PatternFill("solid", fgColor="FBDCC3")
XL_DONE = PatternFill("solid", fgColor="8FD3AE")
NO_FILL = PatternFill(fill_type=None)

st.markdown(f"""
<style>
html, body, [class*="css"] {{font-family:{FONT};}}
.block-container {{padding-top:1.3rem;}}
h1 {{color:{TEXT};font-weight:700;letter-spacing:-.01em;}}
.kpi {{background:#FFFFFF;border:1px solid #ECEFF3;border-radius:14px;
      padding:14px 18px;margin-bottom:8px;box-shadow:0 1px 3px rgba(16,24,40,.05);}}
.kpi .lbl {{font-size:.72rem;color:#6B7280;text-transform:uppercase;letter-spacing:.07em;}}
.kpi .val {{font-size:1.15rem;font-weight:700;color:{TEXT};margin-top:2px;}}
.kpi .bar {{height:3px;width:36px;background:{ACCENT};border-radius:3px;margin-bottom:8px;}}
[data-testid="stPlotlyChart"] {{background:#FFFFFF;border:1px solid #ECEFF3;
      border-radius:16px;padding:8px 6px;box-shadow:0 2px 10px rgba(16,24,40,.06);}}
.legend {{font-size:.85rem;color:#4B5563;margin-top:6px;}}
.chip {{display:inline-block;width:22px;height:8px;border-radius:8px;
      vertical-align:middle;margin:0 6px 0 16px;}}
</style>
""", unsafe_allow_html=True)

WEEK_OPTIONS = {"2-Week Lookahead": 2, "3-Week Lookahead": 3,
                "4-Week Lookahead": 4, "6-Week Lookahead": 6}
DAY_MS = 86400000


# ============================================================== helpers
def find_cell(ws, text, max_row=15):
    text = text.upper()
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper().startswith(text):
                return c.row, c.column
    return None


def value_pos_right_of(ws, label):
    pos = find_cell(ws, label)
    if not pos:
        return None, None
    r, c = pos
    for cc in range(c + 1, c + 5):
        v = ws.cell(r, cc).value
        if v not in (None, ""):
            return v, (r, cc)
    return None, None


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def is_marked(v):
    return isinstance(v, str) and v.strip().lower().startswith("x")


def parse_pct(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v * 100 if v <= 1 else v)
    m = re.search(r"(\d+(\.\d+)?)\s*%", str(v))
    return float(m.group(1)) if m else None


@st.cache_data(show_spinner=False)
def parse_workbook(file_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    ws = wb["6wk Schedule"] if "6wk Schedule" in wb.sheetnames else wb.worksheets[0]

    hdr = find_cell(ws, "WORK DESCRIPTION")
    if not hdr:
        raise ValueError("Could not find 'WORK DESCRIPTION'. Please use the lookahead template.")
    hdr_row, desc_col = hdr
    status_pos = find_cell(ws, "STATUS")
    if not status_pos:
        raise ValueError("Could not find the 'STATUS / % COMPLETE' column.")
    status_col = status_pos[1]
    dur_pos = find_cell(ws, "DURATION")
    day_start = (dur_pos[1] + 1) if dur_pos else desc_col + 1

    info = {
        "title": ws.cell(hdr_row - 3, desc_col).value or "",
        "project": value_pos_right_of(ws, "PROJECT NAME")[0] or "",
        "number": str(value_pos_right_of(ws, "PROJECT NUMBER")[0] or ""),
        "super": value_pos_right_of(ws, "SUPERINTENDENT")[0] or "",
        "from": to_date(value_pos_right_of(ws, "FROM")[0]),
        "sheet": ws.title, "hdr_row": hdr_row, "desc_col": desc_col,
        "id_col": desc_col - 1, "day_cols": list(range(day_start, status_col)),
        "status_col": status_col, "man_col": status_col + 1, "notes_col": status_col + 2,
    }
    if not info["from"]:
        raise ValueError("The 'FROM:' date is missing or not a valid date.")

    rows, section = [], "General"
    # section headers are the bold rows in the template; fall back to "no x marks"
    use_bold = any(ws.cell(r, desc_col).font.b and ws.cell(r, desc_col).value
                   for r in range(hdr_row + 2, ws.max_row + 1))
    for r in range(hdr_row + 2, ws.max_row + 1):
        desc = ws.cell(r, desc_col).value
        if desc is None or str(desc).strip() in ("", "0"):
            continue
        desc = str(desc).strip()
        offsets = [i for i, c in enumerate(info["day_cols"]) if is_marked(ws.cell(r, c).value)]
        bold = bool(ws.cell(r, desc_col).font and ws.cell(r, desc_col).font.b)
        if bold or (not use_bold and not offsets):
            section = desc            # trade / area header row (bold in template)
            continue
        if not offsets:               # activity with nothing scheduled -> not shown
            continue
        rows.append({
            "row": r, "Section": section, "Activity": desc,
            "Sched ID": ws.cell(r, info["id_col"]).value or "",
            "offsets": offsets, "progress": parse_pct(ws.cell(r, status_col).value),
            "Manpower": ws.cell(r, info["man_col"]).value or "",
            "Notes": ws.cell(r, info["notes_col"]).value or "",
        })
    return info, rows


def done_offsets(offsets, pct):
    """First N planned days (chronological) count as complete."""
    if pct is None or pct <= 0:
        return set()
    return set(offsets[:round(len(offsets) * min(pct, 100) / 100)])


def segments(offs):
    segs, cur = [], [offs[0]]
    for o in offs[1:]:
        if o == cur[-1] + 1:
            cur.append(o)
        else:
            segs.append(cur)
            cur = [o]
    segs.append(cur)
    return segs


def short(text, n=42):
    return text if len(text) <= n else text[: n - 1] + "…"


# ============================================================== sleek gantt
def build_gantt(info, rows, progress, n_days, n_weeks, show_labels=True, compact=False):
    start = info["from"]
    win_end = start + timedelta(days=n_days)
    sections = list(dict.fromkeys(a["Section"] for a in rows))
    colors, i = {}, 0
    for s in sections:
        if is_bird_section(s):
            colors[s] = BIRD_COLORS
        else:
            colors[s] = PALETTE[i % len(PALETTE)]
            i += 1

    # y categories: section header rows + activity rows (unique keys)
    cats, ticktext, sec_rows = [], [], []
    visible = [a for a in rows if any(o < n_days for o in a["offsets"])]
    if not visible:
        return None
    for s in sections:
        acts = [a for a in visible if a["Section"] == s]
        if not acts:
            continue
        key = f"§{s}"
        cats.append(key)
        sec_rows.append((key, s))
        ticktext.append(f"<span style='color:{colors[s][1]}'><b>{s.upper()}</b></span>")
        for a in acts:
            a_key = f"r{a['row']}"
            cats.append(a_key)
            ticktext.append(f"<span style='color:{TEXT}'>{short(a['Activity'])}</span>")

    fig = go.Figure()
    bar_w = 0.5 if compact else 0.62
    shown = set()

    for a in visible:
        key = f"r{a['row']}"
        light, dark = colors[a["Section"]]
        offs = [o for o in a["offsets"] if o < n_days]
        pct = progress.get(a["row"])
        done = done_offsets(a["offsets"], pct)
        pct_txt = f"{pct:.0f}%" if pct is not None else "—"
        for seg in segments(offs):
            s0 = start + timedelta(days=seg[0])
            s1 = start + timedelta(days=seg[-1])
            hover = (f"<b>{a['Activity']}</b><br>"
                     f"<span style='color:{MUTED}'>{a['Section']}</span><br><br>"
                     f"📅 {s0:%a %b %d} → {s1:%a %b %d}  ·  {len(seg)} d<br>"
                     f"✅ Progress: <b>{pct_txt}</b>"
                     + (f"<br>👷 Manpower: {a['Manpower']}" if a["Manpower"] else ""))
            fig.add_trace(go.Bar(
                y=[key], x=[len(seg) * DAY_MS - DAY_MS * 0.08], base=[s0],
                orientation="h", width=bar_w,
                marker=dict(color=light, cornerradius=6, line=dict(width=0)),
                name=a["Section"], legendgroup=a["Section"],
                showlegend=a["Section"] not in shown,
                hovertemplate=hover + "<extra></extra>"))
            shown.add(a["Section"])
            d = [o for o in seg if o in done]
            for dseg in (segments(d) if d else []):
                fig.add_trace(go.Bar(
                    y=[key], x=[len(dseg) * DAY_MS - DAY_MS * 0.08],
                    base=[start + timedelta(days=dseg[0])],
                    orientation="h", width=bar_w,
                    marker=dict(color=dark, cornerradius=6, line=dict(width=0)),
                    legendgroup=a["Section"], showlegend=False,
                    hovertemplate=hover + "<extra></extra>"))
        if show_labels and pct is not None:
            complete = pct >= 100
            fig.add_annotation(
                x=start + timedelta(days=offs[-1] + 1), y=key, xanchor="left", xshift=6,
                text=f"<b>{'✓ ' if complete else ''}{pct:.0f}%</b>", showarrow=False,
                font=dict(size=11, color="#FFFFFF"),
                bgcolor="#22A06B" if complete else dark, borderpad=3, opacity=0.95)

    # week bands (alternating) + week labels on top
    for w in range(n_weeks):
        w0 = start + timedelta(days=7 * w)
        w1 = min(w0 + timedelta(days=7), win_end)
        if w0 >= win_end:
            break
        if w % 2 == 1:
            fig.add_vrect(x0=w0, x1=w1, fillcolor="#F8FAFC", opacity=1,
                          line_width=0, layer="below")
        label = "THIS WEEK" if w == 0 else ("NEXT WEEK" if w == 1 else f"WEEK {w + 1}")
        fig.add_annotation(x=w0 + (w1 - w0) / 2, y=1.0, yref="paper", yanchor="bottom",
                           yshift=34, showarrow=False,
                           text=f"<b>{label}</b>  <span style='color:{MUTED}'>{w0:%b %d}</span>",
                           font=dict(size=11, color="#4B5563"))
        if w > 0:
            fig.add_vline(x=w0, line_width=1, line_color="#E5E7EB", layer="below")

    # weekends - very subtle hatch-free tint
    for i in range(n_days):
        d = start + timedelta(days=i)
        if d.weekday() >= 5:
            fig.add_vrect(x0=d, x1=d + timedelta(days=1), fillcolor="#EEF1F5",
                          opacity=0.6, line_width=0, layer="below")

    # section separator lines
    for key, _ in sec_rows[1:]:
        fig.add_shape(type="line", xref="paper", x0=0, x1=1, y0=key, y1=key,
                      yref="y", line=dict(color="#E5E7EB", width=1), layer="below")

    # today marker
    today = date.today()
    if start <= today < win_end:
        t = datetime.combine(today, datetime.min.time()) + timedelta(hours=12)
        fig.add_vline(x=t, line_color=TODAY, line_width=1.5, line_dash="dot")
        fig.add_annotation(x=t, y=0, yref="paper", yanchor="top", yshift=-4,
                           text="<b>TODAY</b>", showarrow=False,
                           font=dict(size=10, color="#FFFFFF"), bgcolor=TODAY, borderpad=3)

    # invisible anchor so every row (incl. section headers) is kept on the axis
    fig.add_trace(go.Scatter(x=[start] * len(cats), y=cats, mode="markers",
                             marker=dict(opacity=0), hoverinfo="skip", showlegend=False))

    row_h = 30 if compact else 38
    fig.update_layout(
        barmode="overlay", bargap=0.35,
        height=max(380, row_h * len(cats) + 170),
        plot_bgcolor="#FFFFFF", paper_bgcolor="#FFFFFF",
        font=dict(family=FONT, color=TEXT, size=12),
        margin=dict(l=10, r=60, t=80, b=40),
        showlegend=False,
        hoverlabel=dict(bgcolor="#FFFFFF", bordercolor="#E5E7EB",
                        font=dict(family=FONT, size=12, color=TEXT)),
    )
    fig.update_yaxes(type="category", categoryorder="array", categoryarray=cats[::-1],
                     range=[-0.6, len(cats) - 0.4],
                     tickvals=cats, ticktext=ticktext, title="",
                     showgrid=False, zeroline=False, ticks="outside", ticklen=8, tickcolor="#FFFFFF", automargin=True)
    fig.update_xaxes(type="date", range=[start, win_end], side="top",
                     dtick=DAY_MS, tickformat="%a<br><b>%d</b>",
                     tickfont=dict(size=10, color="#6B7280"),
                     tick0=datetime.combine(start, datetime.min.time()) + timedelta(hours=12),
                     showgrid=False, gridcolor=GRID,
                     zeroline=False, showline=False, ticks="")
    return fig, colors


# ============================================================== excel export
def export_excel(file_bytes, info, rows, progress, n_weeks, is_xlsm):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), keep_vba=is_xlsm)
    ws = wb[info["sheet"]]
    n_days = min(n_weeks * 7, len(info["day_cols"]))
    start, hdr_row = info["from"], info["hdr_row"]

    _, to_pos = value_pos_right_of(ws, "TO:")
    if to_pos:
        c = ws.cell(*to_pos)
        c.value = start + timedelta(days=n_days - 1)
        c.number_format = "mm/dd/yyyy"

    for i, col in enumerate(info["day_cols"]):
        d = start + timedelta(days=i)
        ws.column_dimensions[get_column_letter(col)].hidden = i >= n_days
        ws.cell(hdr_row, col).value = "MTWTFSS"[d.weekday()]
        ws.cell(hdr_row + 1, col).value = d.day
        ws.cell(hdr_row + 1, col).number_format = "0"

    for a in rows:
        pct = progress.get(a["row"])
        done = done_offsets(a["offsets"], pct)
        for i, col in enumerate(info["day_cols"]):
            cell = ws.cell(a["row"], col)
            if i >= n_days:
                if is_marked(cell.value):
                    cell.value, cell.fill = None, NO_FILL
                continue
            if i in a["offsets"]:
                cell.fill = XL_DONE if i in done else XL_PLANNED
                cell.alignment = Alignment(horizontal="center")
        if pct is not None:                       # blank = leave as-is
            sc = ws.cell(a["row"], info["status_col"])
            sc.value = pct / 100
            sc.number_format = "0%"
            sc.alignment = Alignment(horizontal="center")
            sc.font = Font(bold=True, color="22A06B" if pct >= 100 else "1F2937")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def kpi(col, label, value):
    col.markdown(f'<div class="kpi"><div class="bar"></div><div class="lbl">{label}</div>'
                 f'<div class="val">{value}</div></div>', unsafe_allow_html=True)


# ============================================================== sidebar
with st.sidebar:
    if LOGO_PATH.exists():
        st.image(str(LOGO_PATH), use_container_width=True)
    else:
        st.caption("⚠️ Bird.png not found. Place it in the same folder as this script.")
    st.divider()
    st.header("📂 1. Upload")
    up = st.file_uploader("Lookahead Excel (6wk Schedule template)", type=["xlsx", "xlsm"])
    st.header("📅 2. Lookahead window")
    choice = st.radio("Show", list(WEEK_OPTIONS), index=3, label_visibility="collapsed")
    n_weeks = WEEK_OPTIONS[choice]
    st.header("🎨 3. Chart display")
    show_labels = st.toggle("Show % labels on bars", value=True)
    compact = st.toggle("Compact rows", value=False)
    st.header("✅ 4. Progress")
    st.caption("Enter % complete in the **Update Progress** tab. "
               "Leave blank if there is no update; blanks are not written to Excel.")
    if st.button("↺ Reset progress to file values", use_container_width=True):
        st.session_state.pop("file_key", None)
        st.rerun()

# ============================================================== main
if not up:
    st.title("🏗️ Construction Lookahead Dashboard")
    st.info("Upload your lookahead Excel file in the sidebar to get started.")
    st.markdown("""
**How it works**
1. Upload the *6wk Schedule* Excel file.
2. Pick a **2, 3, 4 or 6-week** lookahead.
3. Enter **% complete** for any activity in the *Update Progress* tab (optional).
4. Review the Gantt chart: the darker part of each bar is the completed portion.
5. **Download** the lookahead in the same Excel format.
""")
    st.stop()

file_bytes = up.getvalue()
try:
    info, rows = parse_workbook(file_bytes)
except Exception as e:
    st.error(f"Could not read the file: {e}")
    st.stop()

file_key = f"{up.name}-{len(file_bytes)}"
if st.session_state.get("file_key") != file_key:
    st.session_state["file_key"] = file_key
    st.session_state.pop("progress_editor", None)
    st.session_state["progress_df"] = pd.DataFrame([{
        "row": a["row"], "Section": a["Section"], "Activity": a["Activity"],
        "Planned Days": len(a["offsets"]), "% Complete": a["progress"],
    } for a in rows])

n_days = min(n_weeks * 7, len(info["day_cols"]))
win_start = info["from"]
win_end = win_start + timedelta(days=n_days - 1)

st.title(f"🏗️ {info['title'] or info['project']} — {choice}")
c1, c2, c3, c4 = st.columns(4)
kpi(c1, "Project Name", f"{info['project']} ({info['number']})")
kpi(c2, "Superintendent", info["super"] or "—")
kpi(c3, "Planned From", f"{win_start:%a, %b %d, %Y}")
kpi(c4, "Lookahead To", f"{win_end:%a, %b %d, %Y}")

tab_g, tab_p, tab_l = st.tabs(["📊 Gantt Chart", "✏️ Update Progress", "📋 Activity List"])

with tab_p:
    st.markdown("Type the **% complete** (0–100) for each activity. "
                "Leave a cell **blank** if there is no progress to report.")
    edited = st.data_editor(
        st.session_state["progress_df"], key="progress_editor",
        hide_index=True, use_container_width=True,
        disabled=["Section", "Activity", "Planned Days"],
        column_order=["Section", "Activity", "Planned Days", "% Complete"],
        column_config={"% Complete": st.column_config.NumberColumn(
            "% Complete", min_value=0, max_value=100, step=5, format="%d %%",
            help="Leave blank if no progress")})
    st.caption("Progress is applied to the earliest planned days first.")

progress = {int(r["row"]): (None if pd.isna(r["% Complete"]) else float(r["% Complete"]))
            for _, r in edited.iterrows()}

in_win = [a for a in rows if any(o < n_days for o in a["offsets"])]
entered = [progress[a["row"]] for a in in_win if progress.get(a["row"]) is not None]
m1, m2, m3, m4 = st.columns(4)
m1.metric("Activities in window", len(in_win))
m2.metric("Progress entered", f"{len(entered)} / {len(in_win)}")
m3.metric("Completed (100%)", sum(1 for p in entered if p >= 100))
m4.metric("Avg. progress (entered)", f"{sum(entered)/len(entered):.0f}%" if entered else "—")

with tab_g:
    result = build_gantt(info, rows, progress, n_days, n_weeks, show_labels, compact)
    if result is None:
        st.warning("No activities are planned within this lookahead window.")
    else:
        fig, colors = result
        st.plotly_chart(fig, use_container_width=True,
                        config={"displaylogo": False,
                                "modeBarButtonsToRemove": ["lasso2d", "select2d"],
                                "toImageButtonOptions": {"filename": f"{info['project']}_{n_weeks}wk_gantt",
                                                         "scale": 2}})
        chips = "".join(f'<span class="chip" style="background:{d}"></span>{s}'
                        for s, (_, d) in colors.items()
                        if any(a["Section"] == s for a in in_win))
        st.markdown(
            f'<div class="legend">{chips}</div><div class="legend">'
            f'<span class="chip" style="background:#C7D7F5"></span>Planned (light)'
            f'<span class="chip" style="background:#4F7CDB"></span>Completed (dark)'
            f'<span class="chip" style="background:#EEF1F5"></span>Weekend'
            f'<span class="chip" style="background:{TODAY}"></span>Today</div>',
            unsafe_allow_html=True)

with tab_l:
    table = []
    for a in in_win:
        offs = [o for o in a["offsets"] if o < n_days]
        table.append({
            "Section": a["Section"], "Sched ID": a["Sched ID"], "Activity": a["Activity"],
            "Start": win_start + timedelta(days=offs[0]),
            "Finish": win_start + timedelta(days=offs[-1]),
            "Duration (days)": len(offs), "% Complete": progress.get(a["row"]),
            "Manpower": a["Manpower"], "Notes": a["Notes"],
        })
    tdf = pd.DataFrame(table)
    if not tdf.empty:
        sec = st.multiselect("Filter by section", sorted(tdf["Section"].unique()))
        if sec:
            tdf = tdf[tdf["Section"].isin(sec)]
    st.dataframe(tdf, use_container_width=True, hide_index=True, column_config={
        "% Complete": st.column_config.ProgressColumn(
            "% Complete", min_value=0, max_value=100, format="%d%%"),
        "Start": st.column_config.DateColumn(format="ddd MMM DD"),
        "Finish": st.column_config.DateColumn(format="ddd MMM DD")})

# ============================================================== download
st.divider()
is_xlsm = up.name.lower().endswith(".xlsm")
ext = "xlsm" if is_xlsm else "xlsx"
out_bytes = export_excel(file_bytes, info, rows, progress, n_weeks, is_xlsm)
fname = f"{info['project'] or 'Project'}_{n_weeks}wk_Lookahead_{win_start:%Y%m%d}.{ext}"
st.download_button(
    f"⬇️ Download {choice} with progress (Excel)", data=out_bytes, file_name=fname,
    mime=("application/vnd.ms-excel.sheet.macroEnabled.12" if is_xlsm else
          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    type="primary", use_container_width=True)
st.caption("Same template layout. Planned days are shaded peach, completed days green, "
           "% complete goes in STATUS / % COMPLETE only where you entered it, "
           "and days outside the window are hidden. -CessDR")
