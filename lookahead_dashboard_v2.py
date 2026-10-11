"""
Construction Lookahead Dashboard v2 (2 / 3 / 4 / 6 week) with Progress Tracking
-------------------------------------------------------------------------------
- Upload the Bird "6wk Schedule" lookahead template (.xlsx / .xlsm)
- Pick a 2, 3, 4 or 6-week window -> Gantt chart
- Enter % complete per activity -> shown directly on the Gantt bars
- Download the lookahead in the SAME Excel format (progress written to
  STATUS / % COMPLETE only where entered; blanks are left as they were)

Run:
    pip install streamlit pandas plotly openpyxl
    streamlit run lookahead_dashboard_v2.py
"""

import io
import re
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

# ============================================================== look & feel
st.set_page_config(page_title="Lookahead Schedule Dashboard",
                   page_icon="🏗️", layout="wide")

# Soft, eye-friendly palette (muted pastels for planned, deeper tone for done)
PALETTE = [
    ("#A8C5E2", "#3F6E9E"),  # soft blue
    ("#B9DDC4", "#3E8E5E"),  # sage green
    ("#F6D3A8", "#C9772B"),  # warm sand / orange
    ("#D7C4E6", "#7A5A9E"),  # lavender
    ("#F2C1C1", "#B85C5C"),  # dusty rose
    ("#BFE3E0", "#3C8A84"),  # teal mist
    ("#E8DFA8", "#9C8A2E"),  # soft olive
    ("#CFD6DE", "#5B6878"),  # slate
]
ACCENT = "#E8833A"           # softened Bird orange
TEXT = "#2B3440"

# Excel colours for the export (same soft theme)
XL_PLANNED = PatternFill("solid", fgColor="F6D3A8")   # planned, not yet done
XL_DONE = PatternFill("solid", fgColor="8CC6A0")      # completed portion
NO_FILL = PatternFill(fill_type=None)

st.markdown(f"""
<style>
.block-container {{padding-top: 1.4rem;}}
h1 {{color:{TEXT};}}
.kpi {{background:#FAFBFC;border:1px solid #E6E9EE;border-left:6px solid {ACCENT};
      border-radius:10px;padding:12px 16px;margin-bottom:8px;}}
.kpi .lbl {{font-size:0.75rem;color:#6B7684;text-transform:uppercase;letter-spacing:.05em;}}
.kpi .val {{font-size:1.2rem;font-weight:700;color:{TEXT};}}
.legend-chip {{display:inline-block;width:14px;height:14px;border-radius:3px;
      vertical-align:middle;margin:0 6px 0 14px;}}
</style>
""", unsafe_allow_html=True)

WEEK_OPTIONS = {"2-Week Lookahead": 2, "3-Week Lookahead": 3,
                "4-Week Lookahead": 4, "6-Week Lookahead": 6}


# ============================================================== helpers
def find_cell(ws, text, max_row=15):
    text = text.upper()
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper().startswith(text):
                return c.row, c.column
    return None


def value_pos_right_of(ws, label):
    pos = find_cell(ws, label)
    if not pos:
        return None, None
    r, c = pos
    for cc in range(c + 1, c + 5):
        v = ws.cell(r, cc).value
        if v not in (None, ""):
            return v, (r, cc)
    return None, None


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def is_marked(v):
    return isinstance(v, str) and v.strip().lower().startswith("x")


def parse_pct(v):
    """Read an existing STATUS / % COMPLETE value -> 0-100, or None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v * 100 if v <= 1 else v)
    m = re.search(r"(\d+(\.\d+)?)\s*%", str(v))
    return float(m.group(1)) if m else None


@st.cache_data(show_spinner=False)
def parse_workbook(file_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    ws = wb["6wk Schedule"] if "6wk Schedule" in wb.sheetnames else wb.worksheets[0]

    hdr = find_cell(ws, "WORK DESCRIPTION")
    if not hdr:
        raise ValueError("Could not find 'WORK DESCRIPTION'. Please use the lookahead template.")
    hdr_row, desc_col = hdr
    status_pos = find_cell(ws, "STATUS")
    if not status_pos:
        raise ValueError("Could not find the 'STATUS / % COMPLETE' column.")
    status_col = status_pos[1]
    dur_pos = find_cell(ws, "DURATION")
    day_start = (dur_pos[1] + 1) if dur_pos else desc_col + 1

    info = {
        "title": ws.cell(hdr_row - 3, desc_col).value or "",
        "project": value_pos_right_of(ws, "PROJECT NAME")[0] or "",
        "number": str(value_pos_right_of(ws, "PROJECT NUMBER")[0] or ""),
        "super": value_pos_right_of(ws, "SUPERINTENDENT")[0] or "",
        "from": to_date(value_pos_right_of(ws, "FROM")[0]),
        "sheet": ws.title, "hdr_row": hdr_row, "desc_col": desc_col,
        "id_col": desc_col - 1, "day_cols": list(range(day_start, status_col)),
        "status_col": status_col, "man_col": status_col + 1, "notes_col": status_col + 2,
    }
    if not info["from"]:
        raise ValueError("The 'FROM:' date is missing or not a valid date.")

    rows, section = [], "General"
    for r in range(hdr_row + 2, ws.max_row + 1):
        desc = ws.cell(r, desc_col).value
        if desc is None or str(desc).strip() in ("", "0"):
            continue
        desc = str(desc).strip()
        offsets = [i for i, c in enumerate(info["day_cols"]) if is_marked(ws.cell(r, c).value)]
        if not offsets:
            section = desc            # trade / area header row
            continue
        status = ws.cell(r, status_col).value
        rows.append({
            "row": r, "Section": section, "Activity": desc,
            "Sched ID": ws.cell(r, info["id_col"]).value or "",
            "offsets": offsets, "status_raw": status,
            "progress": parse_pct(status),
            "Manpower": ws.cell(r, info["man_col"]).value or "",
            "Notes": ws.cell(r, info["notes_col"]).value or "",
        })
    return info, rows


def done_offsets(offsets, pct):
    """First N planned days (chronological) count as complete."""
    if pct is None or pct <= 0:
        return set()
    n = round(len(offsets) * min(pct, 100) / 100)
    return set(offsets[:n])


def segments(offs):
    segs, cur = [], [offs[0]]
    for o in offs[1:]:
        if o == cur[-1] + 1:
            cur.append(o)
        else:
            segs.append(cur); cur = [o]
    segs.append(cur)
    return segs


def build_gantt(info, rows, progress, n_days, n_weeks):
    start = info["from"]
    sections = list(dict.fromkeys(a["Section"] for a in rows))
    colors = {s: PALETTE[i % len(PALETTE)] for i, s in enumerate(sections)}
    fig = go.Figure()
    y_labels, shown_legend = [], set()
    DAY = 86400000

    for a in rows:
        offs = [o for o in a["offsets"] if o < n_days]
        if not offs:
            continue
        label = a["Activity"]
        y_labels.append(label)
        pct = progress.get(a["row"])
        done = done_offsets(a["offsets"], pct)
        light, dark = colors[a["Section"]]
        pct_txt = f"{pct:.0f}%" if pct is not None else "No progress entered"

        for seg in segments(offs):
            s_date = start + timedelta(days=seg[0])
            hover = (f"<b>{label}</b><br>{a['Section']}<br>"
                     f"{s_date:%a %b %d} → {start + timedelta(days=seg[-1]):%a %b %d}"
                     f"<br>Days: {len(seg)}<br>Progress: {pct_txt}"
                     + (f"<br>Manpower: {a['Manpower']}" if a["Manpower"] else ""))
            # planned (light) bar
            fig.add_trace(go.Bar(
                y=[label], x=[len(seg) * DAY], base=[s_date], orientation="h",
                marker=dict(color=light, line=dict(color=dark, width=1)),
                name=a["Section"], legendgroup=a["Section"],
                showlegend=a["Section"] not in shown_legend,
                hovertemplate=hover + "<extra></extra>", width=0.6))
            shown_legend.add(a["Section"])
            # completed (dark) overlay, per contiguous done block
            d = [o for o in seg if o in done]
            if d:
                for dseg in segments(d):
                    fig.add_trace(go.Bar(
                        y=[label], x=[len(dseg) * DAY],
                        base=[start + timedelta(days=dseg[0])], orientation="h",
                        marker=dict(color=dark), width=0.6, showlegend=False,
                        legendgroup=a["Section"],
                        hovertemplate=hover + "<extra></extra>"))
        # % label right of the last bar
        if pct is not None:
            fig.add_annotation(x=start + timedelta(days=offs[-1] + 1), y=label,
                               text=f"<b>{pct:.0f}%</b>", showarrow=False,
                               xanchor="left", xshift=6,
                               font=dict(size=12, color="#2E7D4F" if pct >= 100 else TEXT))

    if not y_labels:
        return None
    win_end = start + timedelta(days=n_days)
    for i in range(n_days):
        d = start + timedelta(days=i)
        if d.weekday() >= 5:
            fig.add_vrect(x0=d, x1=d + timedelta(days=1), fillcolor="#EEF0F3",
                          opacity=0.7, line_width=0, layer="below")
    for w in range(1, n_weeks):
        fig.add_vline(x=start + timedelta(days=7 * w), line_dash="dot",
                      line_color=ACCENT, opacity=0.6)
    today = date.today()
    if start <= today < win_end:
        fig.add_vline(x=today, line_color="#D9534F", line_width=2)

    order = list(dict.fromkeys(y_labels))
    fig.update_layout(
        barmode="overlay", height=max(380, 44 * len(order) + 150),
        plot_bgcolor="white", paper_bgcolor="white",
        font=dict(color=TEXT, size=13),
        margin=dict(l=10, r=40, t=60, b=10),
        legend=dict(orientation="h", y=-0.06, title=""),
        hoverlabel=dict(bgcolor="white", font_size=12))
    fig.update_yaxes(categoryorder="array", categoryarray=order[::-1], title="",
                     gridcolor="#F1F3F6")
    fig.update_xaxes(type="date", range=[start, win_end + timedelta(days=1)],
                     dtick=DAY, tickformat="%a<br>%b %d", side="top",
                     showgrid=True, gridcolor="#E6E9EE")
    return fig


def export_excel(file_bytes, info, rows, progress, n_weeks, is_xlsm):
    """Same template, trimmed to the window, progress written where entered."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), keep_vba=is_xlsm)
    ws = wb[info["sheet"]]
    n_days = min(n_weeks * 7, len(info["day_cols"]))
    start, hdr_row = info["from"], info["hdr_row"]

    _, to_pos = value_pos_right_of(ws, "TO:")
    if to_pos:
        c = ws.cell(*to_pos)
        c.value = start + timedelta(days=n_days - 1)
        c.number_format = "mm/dd/yyyy"

    for i, col in enumerate(info["day_cols"]):
        d = start + timedelta(days=i)
        in_win = i < n_days
        ws.column_dimensions[get_column_letter(col)].hidden = not in_win
        ws.cell(hdr_row, col).value = "MTWTFSS"[d.weekday()]
        ws.cell(hdr_row + 1, col).value = d.day
        ws.cell(hdr_row + 1, col).number_format = "0"

    for a in rows:
        pct = progress.get(a["row"])
        done = done_offsets(a["offsets"], pct)
        for i, col in enumerate(info["day_cols"]):
            cell = ws.cell(a["row"], col)
            if i >= n_days:
                if is_marked(cell.value):
                    cell.value, cell.fill = None, NO_FILL
                continue
            if i in a["offsets"]:
                cell.fill = XL_DONE if i in done else XL_PLANNED
                cell.alignment = Alignment(horizontal="center")
        # write progress only if the user entered one; otherwise leave as-is
        if pct is not None:
            sc = ws.cell(a["row"], info["status_col"])
            sc.value = pct / 100
            sc.number_format = "0%"
            sc.alignment = Alignment(horizontal="center")
            sc.font = Font(bold=True, color="2E7D4F" if pct >= 100 else "2B3440")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def kpi(col, label, value):
    col.markdown(f'<div class="kpi"><div class="lbl">{label}</div>'
                 f'<div class="val">{value}</div></div>', unsafe_allow_html=True)


# ============================================================== sidebar
with st.sidebar:
    st.header("📂 1. Upload")
    up = st.file_uploader("Lookahead Excel (6wk Schedule template)", type=["xlsx", "xlsm"])
    st.header("📅 2. Lookahead window")
    choice = st.radio("Show", list(WEEK_OPTIONS), index=3, label_visibility="collapsed")
    n_weeks = WEEK_OPTIONS[choice]
    st.header("✅ 3. Progress")
    st.caption("Enter % complete in the **Update Progress** tab. "
               "Leave blank if there is no update; blanks are not written to Excel.")
    if st.button("↺ Reset progress to file values", use_container_width=True):
        st.session_state.pop("progress_editor", None)
        st.session_state.pop("progress_df", None)
        st.rerun()

# ============================================================== main
if not up:
    st.title("🏗️ Construction Lookahead Dashboard")
    st.info("Upload your lookahead Excel file in the sidebar to get started.")
    st.markdown("""
**How it works**
1. Upload the *6wk Schedule* Excel file.
2. Pick a **2, 3, 4 or 6-week** lookahead.
3. Enter **% complete** for any activity in the *Update Progress* tab (optional).
4. Review the Gantt chart: the darker part of each bar is the completed portion.
5. **Download** the lookahead in the same Excel format.
""")
    st.stop()

file_bytes = up.getvalue()
try:
    info, rows = parse_workbook(file_bytes)
except Exception as e:
    st.error(f"Could not read the file: {e}")
    st.stop()

# Progress table kept in session (reset when a new file is uploaded)
file_key = f"{up.name}-{len(file_bytes)}"
if st.session_state.get("file_key") != file_key:
    st.session_state["file_key"] = file_key
    st.session_state.pop("progress_editor", None)
    st.session_state["progress_df"] = pd.DataFrame([{
        "row": a["row"], "Section": a["Section"], "Activity": a["Activity"],
        "Planned Days": len(a["offsets"]),
        "% Complete": a["progress"],
    } for a in rows])

n_days = min(n_weeks * 7, len(info["day_cols"]))
win_start = info["from"]
win_end = win_start + timedelta(days=n_days - 1)

st.title(f"🏗️ {info['title'] or info['project']} — {choice}")
c1, c2, c3, c4 = st.columns(4)
kpi(c1, "Project Name", f"{info['project']} ({info['number']})")
kpi(c2, "Superintendent", info["super"] or "—")
kpi(c3, "Planned From", f"{win_start:%a, %b %d, %Y}")
kpi(c4, "Lookahead To", f"{win_end:%a, %b %d, %Y}")

tab_g, tab_p, tab_l = st.tabs(["📊 Gantt Chart", "✏️ Update Progress", "📋 Activity List"])

# ---- progress editor (rendered first so the chart uses the latest values)
with tab_p:
    st.markdown("Type the **% complete** (0–100) for each activity. "
                "Leave a cell **blank** if there is no progress to report.")
    edited = st.data_editor(
        st.session_state["progress_df"],
        key="progress_editor", hide_index=True, use_container_width=True,
        disabled=["Section", "Activity", "Planned Days"],
        column_order=["Section", "Activity", "Planned Days", "% Complete"],
        column_config={
            "% Complete": st.column_config.NumberColumn(
                "% Complete", min_value=0, max_value=100, step=5, format="%d %%",
                help="Leave blank if no progress"),
        })
    st.caption("Tip: the darker part of each Gantt bar shows the completed days. "
               "Progress is applied to the earliest planned days first.")

progress = {int(r["row"]): (None if pd.isna(r["% Complete"]) else float(r["% Complete"]))
            for _, r in edited.iterrows()}

# ---- KPIs
in_win = [a for a in rows if any(o < n_days for o in a["offsets"])]
entered = [progress[a["row"]] for a in in_win if progress.get(a["row"]) is not None]
m1, m2, m3, m4 = st.columns(4)
m1.metric("Activities in window", len(in_win))
m2.metric("Progress entered", f"{len(entered)} / {len(in_win)}")
m3.metric("Completed (100%)", sum(1 for p in entered if p >= 100))
m4.metric("Avg. progress (entered)", f"{sum(entered)/len(entered):.0f}%" if entered else "—")

with tab_g:
    fig = build_gantt(info, rows, progress, n_days, n_weeks)
    if fig is None:
        st.warning("No activities are planned within this lookahead window.")
    else:
        st.plotly_chart(fig, use_container_width=True)
        st.markdown(
            f'<span class="legend-chip" style="background:#A8C5E2"></span>Planned'
            f'<span class="legend-chip" style="background:#3F6E9E"></span>Completed'
            f'<span class="legend-chip" style="background:#EEF0F3"></span>Weekend'
            f'<span class="legend-chip" style="background:{ACCENT}"></span>Week break'
            f'<span class="legend-chip" style="background:#D9534F"></span>Today',
            unsafe_allow_html=True)

with tab_l:
    table = []
    for a in in_win:
        offs = [o for o in a["offsets"] if o < n_days]
        p = progress.get(a["row"])
        table.append({
            "Section": a["Section"], "Sched ID": a["Sched ID"], "Activity": a["Activity"],
            "Start": win_start + timedelta(days=offs[0]),
            "Finish": win_start + timedelta(days=offs[-1]),
            "Duration (days)": len(offs),
            "% Complete": p,
            "Manpower": a["Manpower"], "Notes": a["Notes"],
        })
    tdf = pd.DataFrame(table)
    if not tdf.empty:
        sec = st.multiselect("Filter by section", sorted(tdf["Section"].unique()))
        if sec:
            tdf = tdf[tdf["Section"].isin(sec)]
    st.dataframe(tdf, use_container_width=True, hide_index=True, column_config={
        "% Complete": st.column_config.ProgressColumn(
            "% Complete", min_value=0, max_value=100, format="%d%%"),
        "Start": st.column_config.DateColumn(format="ddd MMM DD"),
        "Finish": st.column_config.DateColumn(format="ddd MMM DD"),
    })

# ============================================================== download
st.divider()
is_xlsm = up.name.lower().endswith(".xlsm")
ext = "xlsm" if is_xlsm else "xlsx"
out_bytes = export_excel(file_bytes, info, rows, progress, n_weeks, is_xlsm)
fname = f"{info['project'] or 'Project'}_{n_weeks}wk_Lookahead_{win_start:%Y%m%d}.{ext}"
st.download_button(
    f"⬇️ Download {choice} with progress (Excel)", data=out_bytes, file_name=fname,
    mime=("application/vnd.ms-excel.sheet.macroEnabled.12" if is_xlsm else
          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    type="primary", use_container_width=True)
st.caption("Same template layout. Planned days are shaded sand, completed days green, "
           "% complete goes in STATUS / % COMPLETE only where you entered it, "
           "and days outside the window are hidden.")
